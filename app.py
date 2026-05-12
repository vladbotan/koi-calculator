"""
Streamlit app for calculating the Protection Mainstreaming KOI indicator
from respondent-level Excel datasets with flexible labels and manual mappings.

Install:
    pip install streamlit pandas openpyxl

Run:
    streamlit run app.py

Core KOI logic used:
    - Each of the 8 mandatory KOI question columns is mapped manually.
    - For each question, the user classifies available answer labels as:
        Positive / Negative / Don't know / No answer.
    - No answer responses are excluded from the denominator.
    - Positive result = Positive / (Total answers - No answer).
    - In practice: Positive / (Positive + Negative + Don't know).
    - Empty and unmapped responses are treated as No answer and excluded from the denominator.
    - Sex/gender and disability disaggregation are optional.
    - Donor filtering supports selecting multiple existing donor values; empty selection means all values.
    - Unique people with at least one negative answer is calculated by row: 1 row = 1 person.
    - Element scores are averages of their two questions: SDH, MEA, ACC, PEM.
    - Final PM KOI is the average of SDH, MEA, ACC, PEM.

Notes:
    - For MEA.2, classify the favourable/positive outcome manually. In the standard
      response set, "Not really" and "Not at all" are counted as the numerator.
    - For ACC.1, classify the positive option manually. Usually "At least one available CRM communication means known" is Positive.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Dict, Iterable, List, Optional, Tuple

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

APP_VERSION = "1.0"
APP_LAST_UPDATED = "2026-05-12"
UNMAPPED_WARNING_THRESHOLD = 0.10
NEGATIVE_PERSON_METRIC_LABEL = "People with >=1 negative KOI answer"

QUESTION_LABELS = {
    "sdh_1": "SDH. 1",
    "sdh_2": "SDH. 2",
    "mea_1": "MEA. 1",
    "mea_2": "MEA. 2",
    "acc_1": "ACC. 1",
    "acc_2": "ACC. 2",
    "pem_1": "PEM. 1",
    "pem_2": "PEM. 2",
}

QUESTION_CUES = {
    "sdh_1": "Safety while accessing assistance",
    "sdh_2": "Respectful treatment by staff",
    "mea_1": "Satisfaction with assistance/service",
    "mea_2": "People excluded from assistance",
    "acc_1": "Ability to make a complaint",
    "acc_2": "Complaints followed up",
    "pem_1": "Views taken into account",
    "pem_2": "Information about available assistance",
}

QUESTION_HELP = {
    "sdh_1": "Did the respondent feel safe travelling to, receiving, and returning from the assistance/service?",
    "sdh_2": "Did staff treat the respondent with respect during the intervention?",
    "mea_1": "Was the respondent satisfied with the assistance or service provided?",
    "mea_2": "Does the respondent know of people needing assistance/services who were excluded? This question is reversed for scoring.",
    "acc_1": "Could the respondent channel a suggestion or lodge a complaint if needed?",
    "acc_2": "Were suggestions or complaints responded to or followed up?",
    "pem_1": "Were the respondent's views taken into account about the assistance received?",
    "pem_2": "Did the respondent feel well informed about the assistance/service available?",
}

QUESTION_COLUMN_HINTS = {
    "sdh_1": [
        "sdh 1", "sdh.1", "sdh1", "safe", "safety", "accessing assistance",
        "travel", "returning", "receiving assistance", "feel safe",
    ],
    "sdh_2": [
        "sdh 2", "sdh.2", "sdh2", "respect", "respectful", "treated",
        "staff treatment", "treated with respect", "dignity",
    ],
    "mea_1": [
        "mea 1", "mea.1", "mea1", "satisfaction", "satisfied",
        "assistance satisfaction", "service satisfaction", "quality",
    ],
    "mea_2": [
        "mea 2", "mea.2", "mea2", "excluded", "exclusion",
        "people excluded", "left out", "not included", "need assistance",
    ],
    "acc_1": [
        "acc 1", "acc.1", "acc1", "complaint", "complaints", "suggestion",
        "feedback", "crm", "complaint mechanism", "make a complaint",
    ],
    "acc_2": [
        "acc 2", "acc.2", "acc2", "follow up", "follow-up", "responded",
        "response to complaint", "complaints followed", "feedback followed",
    ],
    "pem_1": [
        "pem 1", "pem.1", "pem1", "views", "opinion", "opinions",
        "taken into account", "participation", "consulted",
    ],
    "pem_2": [
        "pem 2", "pem.2", "pem2", "information", "informed",
        "available assistance", "available service", "know about assistance",
    ],
}

QUESTION_ORDER = ["sdh_1", "sdh_2", "mea_1", "mea_2", "acc_1", "acc_2", "pem_1", "pem_2"]
QUESTION_TO_ELEMENT = {
    "sdh_1": "SDH", "sdh_2": "SDH",
    "mea_1": "MEA", "mea_2": "MEA",
    "acc_1": "ACC", "acc_2": "ACC",
    "pem_1": "PEM", "pem_2": "PEM",
}
ELEMENTS = ["SDH", "MEA", "ACC", "PEM"]

DEFAULT_AGE_BANDS = "5-17, 18-49, 50+"
TEMPLATE_COLUMNS = ["sex", "age", "disability", "donor", *QUESTION_LABELS.values()]
DEFAULT_FEMALE_HINTS = ["female", "f", "woman", "women", "girl", "girls", "feminin", "femeie"]
DEFAULT_MALE_HINTS = ["male", "m", "man", "men", "boy", "boys", "masculin", "barbat"]
DEFAULT_DISABILITY_YES_HINTS = ["yes", "da", "true", "1", "some difficulty", "a lot of difficulty", "cannot do at all"]
DEFAULT_NO_ANSWER_HINTS = [
    "no answer", "do not want to answer", "don't want to answer",
    "refuse", "refused", "prefer not", "n/a", "nan", "blank"
]
DEFAULT_DONT_KNOW_HINTS = [
    "don't know", "dont know", "do not know", "dk", "i don't know", "i dont know",
    "unknown", "not sure", "unsure", "nu stiu", "nu stiu/nu raspund"
]
DEFAULT_POSITIVE_HINTS = [
    "yes completely", "yes, completely", "completely yes", "mostly yes",
    "yes mostly", "mostly", "completely"
]
DEFAULT_NEGATIVE_HINTS = [
    "not really", "not at all", "rather no", "mostly no", "yes a lot", "yes a few"
]

SAMPLE_DATA = [
    {
        "sex": "Female",
        "age": 28,
        "disability": "No",
        "donor": "Program A",
        "SDH. 1": "Yes completely",
        "SDH. 2": "Mostly yes",
        "MEA. 1": "Mostly yes",
        "MEA. 2": "Not at all",
        "ACC. 1": "At least one available CRM communication means known",
        "ACC. 2": "Mostly yes",
        "PEM. 1": "Yes completely",
        "PEM. 2": "Mostly yes",
    },
    {
        "sex": "Male",
        "age": 56,
        "disability": "Yes",
        "donor": "Program A",
        "SDH. 1": "Not really",
        "SDH. 2": "Don't know",
        "MEA. 1": "Not at all",
        "MEA. 2": "Yes completely",
        "ACC. 1": "CRM communication mean(s) mentioned are not available",
        "ACC. 2": "Not really",
        "PEM. 1": "Mostly yes",
        "PEM. 2": "Not really",
    },
    {
        "sex": "Female",
        "age": 16,
        "disability": "No",
        "donor": "Program B",
        "SDH. 1": "Mostly yes",
        "SDH. 2": "No answer",
        "MEA. 1": "Mostly yes",
        "MEA. 2": "Not really",
        "ACC. 1": "At least one available CRM communication means known",
        "ACC. 2": "Don't know",
        "PEM. 1": "Not really",
        "PEM. 2": "Mostly yes",
    },
    {
        "sex": "Male",
        "age": 34,
        "disability": "No",
        "donor": "Program B",
        "SDH. 1": "Mostly yes",
        "SDH. 2": "Mostly yes",
        "MEA. 1": "Don't know",
        "MEA. 2": "Not really",
        "ACC. 1": "At least one available CRM communication means known",
        "ACC. 2": "Yes completely",
        "PEM. 1": "Yes completely",
        "PEM. 2": "Mostly yes",
    },
    {
        "sex": "Female",
        "age": 67,
        "disability": "Yes",
        "donor": "Program C",
        "SDH. 1": "Not at all",
        "SDH. 2": "Not really",
        "MEA. 1": "Mostly yes",
        "MEA. 2": "Mostly yes",
        "ACC. 1": "CRM communication mean(s) mentioned are not available",
        "ACC. 2": "Not at all",
        "PEM. 1": "Don't know",
        "PEM. 2": "Not really",
    },
    {
        "sex": "Male",
        "age": 12,
        "disability": "No",
        "donor": "Program C",
        "SDH. 1": "No answer",
        "SDH. 2": "Mostly yes",
        "MEA. 1": "Yes completely",
        "MEA. 2": "Not at all",
        "ACC. 1": "At least one available CRM communication means known",
        "ACC. 2": "Mostly yes",
        "PEM. 1": "Mostly yes",
        "PEM. 2": "No answer",
    },
    {
        "sex": "Female",
        "age": 43,
        "disability": "No",
        "donor": "Program A",
        "SDH. 1": "Yes completely",
        "SDH. 2": "Yes completely",
        "MEA. 1": "Mostly yes",
        "MEA. 2": "Don't know",
        "ACC. 1": "At least one available CRM communication means known",
        "ACC. 2": "Don't know",
        "PEM. 1": "Mostly yes",
        "PEM. 2": "Yes completely",
    },
    {
        "sex": "Male",
        "age": 51,
        "disability": "Yes",
        "donor": "Program B",
        "SDH. 1": "Mostly yes",
        "SDH. 2": "Not really",
        "MEA. 1": "Not really",
        "MEA. 2": "Yes completely",
        "ACC. 1": "CRM communication mean(s) mentioned are not available",
        "ACC. 2": "Not really",
        "PEM. 1": "Not at all",
        "PEM. 2": "Don't know",
    },
    {
        "sex": "Female",
        "age": 22,
        "disability": "No",
        "donor": "Program C",
        "SDH. 1": "Mostly yes",
        "SDH. 2": "Mostly yes",
        "MEA. 1": "Yes completely",
        "MEA. 2": "Not really",
        "ACC. 1": "At least one available CRM communication means known",
        "ACC. 2": "Yes completely",
        "PEM. 1": "Mostly yes",
        "PEM. 2": "Mostly yes",
    },
]


def clean_text(value: object) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip().lower()
    text = text.replace("\u2019", "'").replace("\u2018", "'").replace("`", "'")
    text = text.replace("\u2013", "-").replace("\u2014", "-")
    text = re.sub(r"[\u200b\u200c\u200d\ufeff]", "", text)
    text = re.sub(r"\s+", " ", text)
    return text


def unique_display_values(series: pd.Series, limit: int = 300) -> List[str]:
    values = []
    seen = set()
    for value in series.dropna().astype(str):
        display = value.strip()
        key = clean_text(display)
        if not key or key in seen:
            continue
        seen.add(key)
        values.append(display)
        if len(values) >= limit:
            break
    return sorted(values, key=lambda x: clean_text(x))


def auto_pick(values: List[str], hints: Iterable[str]) -> List[str]:
    hints_clean = [clean_text(h) for h in hints]
    picked = []
    for value in values:
        v = clean_text(value)
        if any(h == v or h in v for h in hints_clean):
            picked.append(value)
    return picked


def parse_age_band_spec(raw: str) -> List[Tuple[str, Optional[float], Optional[float]]]:
    """Parse bands like '5-17, 18-49, 50+' into label/min/max."""
    bands = []
    if not raw or not raw.strip():
        raise ValueError("Age bands cannot be blank when age disaggregation is enabled.")
    for part in [p.strip() for p in raw.split(",") if p.strip()]:
        normalized = part.replace(" ", "")
        try:
            if normalized.endswith("+"):
                start = float(normalized[:-1])
                bands.append((part, start, None))
            elif "-" in normalized:
                start, end = normalized.split("-", 1)
                start_value = float(start)
                end_value = float(end)
                if end_value < start_value:
                    raise ValueError
                bands.append((part, start_value, end_value))
            else:
                value = float(normalized)
                bands.append((part, value, value))
        except ValueError as exc:
            raise ValueError(f"Invalid age band '{part}'. Use formats like 5-17, 18-49, 50+.") from exc
    if not bands:
        raise ValueError("At least one age band is required when age disaggregation is enabled.")
    return bands


def age_to_band(value: object, bands: List[Tuple[str, Optional[float], Optional[float]]]) -> Optional[str]:
    text = clean_text(value)
    if not text:
        return None

    number = pd.to_numeric(text, errors="coerce")
    if pd.notna(number):
        age = float(number)
        for label, start, end in bands:
            if start is not None and age < start:
                continue
            if end is not None and age > end:
                continue
            return label
        return None

    compact = text.replace(" ", "")
    for label, _, _ in bands:
        label_clean = clean_text(label).replace(" ", "")
        if compact == label_clean or label_clean in compact:
            return label
    return None


def make_mask_from_values(series: pd.Series, selected_values: Iterable[str]) -> pd.Series:
    selected = {clean_text(v) for v in selected_values}
    return series.map(clean_text).isin(selected)


def safe_ratio(numerator: float, denominator: float) -> Optional[float]:
    if denominator == 0:
        return None
    return numerator / denominator


def average(values: Iterable[Optional[float]]) -> Optional[float]:
    valid = [v for v in values if v is not None and pd.notna(v)]
    if not valid:
        return None
    return sum(valid) / len(valid)


@dataclass
class GroupSpec:
    key: str
    label: str
    mask: pd.Series


def apply_donor_filter(df: pd.DataFrame, donor_col: Optional[str], donor_values: Optional[List[str]]) -> pd.DataFrame:
    if not donor_col or not donor_values:
        return df.copy()
    selected = {clean_text(v) for v in donor_values if clean_text(v)}
    if not selected:
        return df.copy()
    return df.loc[df[donor_col].map(clean_text).isin(selected)].copy()


def classify_response(
    value: object,
    positive: Iterable[str],
    negative: Iterable[str],
    dont_know: Iterable[str],
    no_answer: Iterable[str],
) -> str:
    raw = clean_text(value)
    if raw == "":
        return "unmapped_no_answer"
    positive_set = {clean_text(v) for v in positive}
    negative_set = {clean_text(v) for v in negative}
    dont_know_set = {clean_text(v) for v in dont_know}
    no_answer_set = {clean_text(v) for v in no_answer}
    if raw in positive_set:
        return "positive"
    if raw in negative_set:
        return "negative"
    if raw in dont_know_set:
        return "dont_know"
    if raw in no_answer_set:
        return "no_answer"
    # Unmapped values are deliberately excluded from the denominator,
    # but kept separately for review instead of being merged invisibly
    # into explicit no-answer responses.
    return "unmapped_no_answer"


def build_groups(df: pd.DataFrame, config: dict) -> List[GroupSpec]:
    sex_col = config["columns"].get("sex")
    age_col = config["columns"].get("age")
    disability_col = config["columns"].get("disability")

    groups: List[GroupSpec] = []
    all_mask = pd.Series(True, index=df.index)

    if sex_col:
        female = make_mask_from_values(df[sex_col], config["sex_values"].get("female", []))
        male = make_mask_from_values(df[sex_col], config["sex_values"].get("male", []))
        groups.append(GroupSpec("female", "Female", female))
        groups.append(GroupSpec("male", "Male", male))
    else:
        female = pd.Series(False, index=df.index)
        male = pd.Series(False, index=df.index)

    age_disagg = config.get("age_disaggregation", "No disaggregation")
    age_group = pd.Series(None, index=df.index, dtype="object")
    if age_col and age_disagg != "No disaggregation":
        bands = parse_age_band_spec(config.get("age_bands", DEFAULT_AGE_BANDS))
        age_group = df[age_col].map(lambda x: age_to_band(x, bands))
        for band_label, _, _ in bands:
            groups.append(GroupSpec(f"age_{clean_text(band_label).replace(' ', '_').replace('+', 'plus')}", f"Age {band_label}", age_group == band_label))

        if sex_col:
            for band_label, _, _ in bands:
                band_mask = age_group == band_label
                groups.append(GroupSpec(f"female_{clean_text(band_label)}", f"Female {band_label}", female & band_mask))
                groups.append(GroupSpec(f"male_{clean_text(band_label)}", f"Male {band_label}", male & band_mask))

    if disability_col:
        disability_yes = make_mask_from_values(df[disability_col], config.get("disability_yes_values", []))
        groups.append(GroupSpec("disability_yes", "Persons with disabilities", disability_yes))
        if sex_col:
            groups.append(GroupSpec("female_disability_yes", "Female with disabilities", female & disability_yes))
            groups.append(GroupSpec("male_disability_yes", "Male with disabilities", male & disability_yes))

    groups.append(GroupSpec("total", "Total", all_mask))
    return groups


def validate_config(df: pd.DataFrame, config: dict) -> None:
    required = [*QUESTION_ORDER]
    if config.get("age_disaggregation") != "No disaggregation":
        required.append("age")

    missing = []
    for key in required:
        col = config["columns"].get(key)
        if not col or col not in df.columns:
            missing.append(key)
    donor_col = config["columns"].get("donor")
    if donor_col and donor_col not in df.columns:
        missing.append("donor")
    if missing:
        raise ValueError("Missing required mapped columns: " + ", ".join(missing))

    if config.get("age_disaggregation") != "No disaggregation":
        parse_age_band_spec(config.get("age_bands", DEFAULT_AGE_BANDS))

    question_columns: Dict[str, List[str]] = {}
    for q in QUESTION_ORDER:
        col = config["columns"].get(q)
        if col:
            question_columns.setdefault(col, []).append(QUESTION_LABELS[q])
    duplicated_question_columns = {
        col: labels for col, labels in question_columns.items() if len(labels) > 1
    }
    if duplicated_question_columns:
        details = [
            f"{col} used for {', '.join(labels)}"
            for col, labels in duplicated_question_columns.items()
        ]
        raise ValueError("Each KOI question must use a different dataset column: " + "; ".join(details))

    for q in QUESTION_ORDER:
        mapping = config["question_mappings"].get(q, {})
        if not mapping.get("positive"):
            raise ValueError(f"{QUESTION_LABELS[q]} has no Positive responses selected.")
        categories = ["positive", "negative", "dont_know", "no_answer"]
        seen: Dict[str, str] = {}
        overlaps = []
        for category in categories:
            for value in mapping.get(category, []):
                key = clean_text(value)
                if not key:
                    continue
                if key in seen and seen[key] != category:
                    overlaps.append(key)
                seen[key] = category
        if overlaps:
            raise ValueError(f"{QUESTION_LABELS[q]} has response labels selected in more than one category: {', '.join(sorted(set(overlaps)))}")


def negative_hints_for_question(question: str) -> List[str]:
    if question == "mea_2":
        return ["yes completely", "mostly yes", "yes a lot", "yes, a lot", "yes a few", "yes, a few", "a lot", "a few"]
    return DEFAULT_NEGATIVE_HINTS


def mapping_selected_keys(mapping: dict) -> set:
    selected = set()
    for category in ["positive", "negative", "dont_know", "no_answer"]:
        selected.update(clean_text(value) for value in mapping.get(category, []))
    return {value for value in selected if value}


def unmapped_count_and_ratio(series: pd.Series, mapping: dict) -> Tuple[int, float]:
    selected = mapping_selected_keys(mapping)
    total = len(series)
    if total == 0:
        return 0, 0.0
    unmapped_count = int((~series.map(clean_text).isin(selected)).sum())
    return unmapped_count, unmapped_count / total


def format_mapping_values(values: Iterable[str]) -> str:
    selected = [str(value) for value in values if str(value).strip()]
    return " | ".join(selected) if selected else "None selected"


def question_display_label(question: str) -> str:
    return f"{QUESTION_LABELS[question]} - {QUESTION_CUES[question]}"


def build_mapping_review_table(df: pd.DataFrame, config: dict) -> pd.DataFrame:
    rows = []
    for q in QUESTION_ORDER:
        col = config["columns"].get(q)
        mapping = config["question_mappings"].get(q, {})
        notes = []
        unmapped_count = 0
        unmapped_ratio = 0.0
        if col and col in df.columns:
            values = unique_display_values(df[col])
            unmapped_count, unmapped_ratio = unmapped_count_and_ratio(df[col], mapping)
            if auto_pick(values, negative_hints_for_question(q)) and not mapping.get("negative"):
                notes.append("negative-looking values not mapped")
            if auto_pick(values, DEFAULT_DONT_KNOW_HINTS) and not mapping.get("dont_know"):
                notes.append("don't know values not mapped")
            if unmapped_ratio >= UNMAPPED_WARNING_THRESHOLD:
                notes.append("high unmapped/blank share")
        else:
            notes.append("missing column")
        if not mapping.get("positive"):
            notes.append("missing positive mapping")

        rows.append({
            "Question": question_display_label(q),
            "Dataset column": col or "Not selected",
            "Positive": format_mapping_values(mapping.get("positive", [])),
            "Negative": format_mapping_values(mapping.get("negative", [])),
            "Don't know": format_mapping_values(mapping.get("dont_know", [])),
            "No answer": format_mapping_values(mapping.get("no_answer", [])),
            "Unmapped/blank": f"{unmapped_count} ({unmapped_ratio:.1%})",
            "Review notes": "; ".join(notes) if notes else "OK",
        })
    return pd.DataFrame(rows)


def build_preflight_checks(df: pd.DataFrame, config: dict) -> Tuple[List[str], List[str]]:
    errors = []
    warnings = []
    try:
        validate_config(df, config)
    except ValueError as exc:
        errors.append(str(exc))

    donor_col = config["columns"].get("donor")
    donor_values = config.get("donor_values", [])
    if donor_col and donor_col in df.columns and donor_values:
        filtered_df = apply_donor_filter(df, donor_col, donor_values)
        if filtered_df.empty:
            errors.append("The selected donor filter leaves zero rows to calculate.")

    if config.get("age_disaggregation") != "No disaggregation":
        age_col = config["columns"].get("age")
        if age_col and age_col in df.columns:
            try:
                bands = parse_age_band_spec(config.get("age_bands", DEFAULT_AGE_BANDS))
                mapped_age_count = int(df[age_col].map(lambda value: age_to_band(value, bands)).notna().sum())
                if mapped_age_count == 0 and len(df) > 0:
                    warnings.append("No rows matched the selected age bands.")
            except ValueError:
                pass

    sex_col = config["columns"].get("sex")
    if sex_col and sex_col in df.columns:
        sex_values = unique_display_values(df[sex_col])
        if sex_values and not config["sex_values"].get("female"):
            warnings.append("No values are mapped as Female.")
        if sex_values and not config["sex_values"].get("male"):
            warnings.append("No values are mapped as Male.")

    for q in QUESTION_ORDER:
        col = config["columns"].get(q)
        if not col or col not in df.columns:
            continue
        values = unique_display_values(df[col])
        mapping = config["question_mappings"].get(q, {})
        if auto_pick(values, negative_hints_for_question(q)) and not mapping.get("negative"):
            warnings.append(f"{QUESTION_LABELS[q]} has negative-looking values, but none are mapped as Negative.")
        if auto_pick(values, DEFAULT_DONT_KNOW_HINTS) and not mapping.get("dont_know"):
            warnings.append(f"{QUESTION_LABELS[q]} has Don't know values, but none are mapped as Don't know.")
        unmapped_count, unmapped_ratio = unmapped_count_and_ratio(df[col], mapping)
        if unmapped_ratio >= UNMAPPED_WARNING_THRESHOLD:
            warnings.append(
                f"{QUESTION_LABELS[q]} has {unmapped_count} unmapped/blank responses "
                f"({unmapped_ratio:.1%}); these will be excluded from the denominator."
            )

    return list(dict.fromkeys(errors)), list(dict.fromkeys(warnings))


def calculate_koi(raw_df: pd.DataFrame, config: dict) -> dict:
    validate_config(raw_df, config)
    df = apply_donor_filter(raw_df, config["columns"].get("donor"), config.get("donor_values", []))
    df = df.copy()

    for q in QUESTION_ORDER:
        col = config["columns"][q]
        mapping = config["question_mappings"][q]
        df[f"__{q}_norm"] = df[col].map(
            lambda value, m=mapping: classify_response(
                value,
                m.get("positive", []),
                m.get("negative", []),
                m.get("dont_know", []),
                m.get("no_answer", []),
            )
        )

    groups = build_groups(df, config)
    group_keys = [g.key for g in groups]
    sample_sizes = {g.key: int(g.mask.sum()) for g in groups}

    question_counts: Dict[str, Dict[str, Dict[str, int]]] = {}
    question_scores: Dict[str, Dict[str, Optional[float]]] = {}

    for q in QUESTION_ORDER:
        norm_col = f"__{q}_norm"
        question_counts[q] = {}
        question_scores[q] = {}
        for g in groups:
            values = df.loc[g.mask, norm_col].value_counts(dropna=False).to_dict()
            counts = {
                "positive": int(values.get("positive", 0)),
                "negative": int(values.get("negative", 0)),
                "dont_know": int(values.get("dont_know", 0)),
                "no_answer": int(values.get("no_answer", 0)),
                "unmapped_no_answer": int(values.get("unmapped_no_answer", 0)),
            }
            denominator = counts["positive"] + counts["negative"] + counts["dont_know"]
            question_counts[q][g.key] = counts
            question_scores[q][g.key] = safe_ratio(counts["positive"], denominator)

    element_scores = {element: {} for element in ELEMENTS}
    for element in ELEMENTS:
        questions = [q for q in QUESTION_ORDER if QUESTION_TO_ELEMENT[q] == element]
        for group_key in group_keys:
            element_scores[element][group_key] = average(question_scores[q][group_key] for q in questions)

    final_scores = {group_key: average(element_scores[element][group_key] for element in ELEMENTS) for group_key in group_keys}

    negative_any_mask = pd.Series(False, index=df.index)
    for q in QUESTION_ORDER:
        negative_any_mask |= df[f"__{q}_norm"].eq("negative")

    # 1 row = 1 person. A person is counted once if they have at least one negative answer
    # across any of the 8 KOI questions, even if they have negative answers in multiple questions.
    unique_people_with_negative = int(negative_any_mask.sum())

    return {
        "cleaned_df": df,
        "original_rows": len(raw_df),
        "filtered_rows": len(df),
        "groups": groups,
        "group_keys": group_keys,
        "sample_sizes": sample_sizes,
        "question_counts": question_counts,
        "question_scores": question_scores,
        "element_scores": element_scores,
        "final_scores": final_scores,
        "unique_people_with_negative": unique_people_with_negative,
        "config": config,
    }


def build_partner_inputs_table(results: dict) -> pd.DataFrame:
    groups = results["groups"]
    group_keys = results["group_keys"]
    rows = []
    rows.append(["Protection Mainstreaming KOI calculation tool"] + [None] * len(group_keys))
    rows.append(["Original rows"] + [results["original_rows"]] + [None] * (len(group_keys) - 1))
    rows.append(["Rows after donor filter"] + [results["filtered_rows"]] + [None] * (len(group_keys) - 1))
    rows.append([NEGATIVE_PERSON_METRIC_LABEL] + [results.get("unique_people_with_negative")] + [None] * (len(group_keys) - 1))
    rows.append(["Disaggregation"] + [g.label for g in groups])
    rows.append(["Sample size"] + [results["sample_sizes"][k] for k in group_keys])
    rows.append([None] + [None] * len(group_keys))

    for q in QUESTION_ORDER:
        rows.append([QUESTION_LABELS[q]] + [None] * len(group_keys))
        counts = results["question_counts"][q]
        rows.append(["Positive"] + [counts[k]["positive"] for k in group_keys])
        rows.append(["Negative"] + [counts[k]["negative"] for k in group_keys])
        rows.append(["Don't know"] + [counts[k]["dont_know"] for k in group_keys])
        rows.append(["No answer total excluded from denominator"] + [counts[k]["no_answer"] + counts[k]["unmapped_no_answer"] for k in group_keys])
        rows.append(["Explicit no answer"] + [counts[k]["no_answer"] for k in group_keys])
        rows.append(["Unmapped/blank treated as no answer"] + [counts[k]["unmapped_no_answer"] for k in group_keys])
        rows.append(["Question denominator: Positive + Negative + Don't know"] + [counts[k]["positive"] + counts[k]["negative"] + counts[k]["dont_know"] for k in group_keys])
        rows.append(["Question result"] + [results["question_scores"][q][k] for k in group_keys])
        rows.append([None] + [None] * len(group_keys))

    rows.append(["Results of the PM KOI survey for this round"] + [results["final_scores"][k] for k in group_keys])
    rows.append([None] + [None] * len(group_keys))
    for element in ELEMENTS:
        rows.append([f"Average % {element}"] + [results["element_scores"][element][k] for k in group_keys])

    return pd.DataFrame(rows, columns=["Indicator / Response"] + [g.label for g in groups])


def build_summary_table(results: dict) -> pd.DataFrame:
    """Human-readable summary with both percentages and raw response counts.

    Don't know is included in the denominator:
    Positive / (Positive + Negative + Don't know).
    No answer and unmapped/blank are excluded.
    """
    rows = [
        ["Donor column", results["config"]["columns"].get("donor") or "Not used"],
        ["Donor filter", ", ".join(results["config"].get("donor_values", [])) or "All values"],
        ["Original rows", results["original_rows"]],
        ["Rows after donor filter", results["filtered_rows"]],
        [NEGATIVE_PERSON_METRIC_LABEL, results.get("unique_people_with_negative")],
        ["Final PM KOI result", results["final_scores"].get("total")],
        [None, None],
    ]

    for element in ELEMENTS:
        rows.append([f"Average % {element}", results["element_scores"][element].get("total")])

    rows.append([None, None])
    rows.append(["QUESTION RESULTS - TOTAL", None])
    for q in QUESTION_ORDER:
        counts = results["question_counts"][q]["total"]
        denominator = counts["positive"] + counts["negative"] + counts["dont_know"]
        excluded = counts["no_answer"] + counts["unmapped_no_answer"]
        rows.extend([
            [f"{QUESTION_LABELS[q]} result", results["question_scores"][q].get("total")],
            [f"{QUESTION_LABELS[q]} positive", counts["positive"]],
            [f"{QUESTION_LABELS[q]} negative", counts["negative"]],
            [f"{QUESTION_LABELS[q]} don't know", counts["dont_know"]],
            [f"{QUESTION_LABELS[q]} denominator: positive + negative + don't know", denominator],
            [f"{QUESTION_LABELS[q]} no answer excluded", excluded],
            [None, None],
        ])

    return pd.DataFrame(rows, columns=["Metric", "Value"])


def build_export_notes_table() -> pd.DataFrame:
    rows = [
        ["Privacy", "Use anonymized datasets only. Remove names, phone numbers, addresses, ID numbers, and other direct identifiers before upload or export."],
        ["Calculation method", "Question result = Positive / (Positive + Negative + Don't know). No answer and unmapped/blank responses are excluded."],
        ["Final result", "The final PM KOI result is the average of SDH, MEA, ACC, and PEM element scores."],
        [NEGATIVE_PERSON_METRIC_LABEL, "Diagnostic count only. It flags each row/person once if they have any negative KOI answer; it is not the inverse of the final score."],
        ["MEA. 2 note", "For the standard response set, Not really and Not at all are the favourable/positive outcomes."],
        ["Calculator version", APP_VERSION],
        ["Last updated", APP_LAST_UPDATED],
    ]
    return pd.DataFrame(rows, columns=["Topic", "Note"])


def build_unmapped_table(results: dict) -> pd.DataFrame:
    df = results["cleaned_df"]
    rows = []
    for q in QUESTION_ORDER:
        col = results["config"]["columns"][q]
        mapping = results["config"]["question_mappings"][q]
        selected = (
            {clean_text(v) for v in mapping.get("positive", [])}
            | {clean_text(v) for v in mapping.get("negative", [])}
            | {clean_text(v) for v in mapping.get("dont_know", [])}
            | {clean_text(v) for v in mapping.get("no_answer", [])}
        )
        raw_series = df[col]
        unselected_mask = ~raw_series.map(clean_text).isin(selected)
        unmapped = raw_series.loc[unselected_mask].fillna("[blank]").astype(str).replace({"": "[blank]"}).value_counts()
        for raw_value, count in unmapped.items():
            rows.append({
                "question": QUESTION_LABELS[q],
                "mapped_column": col,
                "raw_response": raw_value,
                "count": int(count),
                "treatment": "Treated as No answer and excluded from denominator",
            })
    if not rows:
        return pd.DataFrame([{
            "question": "No unselected responses found",
            "mapped_column": None,
            "raw_response": None,
            "count": None,
            "treatment": None,
        }])
    return pd.DataFrame(rows)

def build_mapping_table(config: dict) -> pd.DataFrame:
    rows = []
    for field, col in config["columns"].items():
        rows.append({"field": field, "selected_column": col or "Not used"})
    for q in QUESTION_ORDER:
        mapping = config["question_mappings"][q]
        rows.append({"field": f"{QUESTION_LABELS[q]} positive", "selected_column": " | ".join(mapping.get("positive", []))})
        rows.append({"field": f"{QUESTION_LABELS[q]} negative", "selected_column": " | ".join(mapping.get("negative", []))})
        rows.append({"field": f"{QUESTION_LABELS[q]} don't know", "selected_column": " | ".join(mapping.get("dont_know", []))})
        rows.append({"field": f"{QUESTION_LABELS[q]} no answer", "selected_column": " | ".join(mapping.get("no_answer", []))})
    return pd.DataFrame(rows)


def format_workbook_bytes(data: bytes) -> bytes:
    buffer = io.BytesIO(data)
    wb = load_workbook(buffer)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for ws in wb.worksheets:
        ws.freeze_panes = "B2"
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = border
                if isinstance(cell.value, float) and 0 <= cell.value <= 1:
                    cell.number_format = "0.00%"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            first = row[0].value
            if isinstance(first, str) and (
                first.startswith(("SDH", "MEA", "ACC", "PEM"))
                or first.startswith("Results of")
                or first.startswith("Average %")
                or first.startswith("Protection Mainstreaming")
            ):
                for cell in row:
                    cell.fill = section_fill
                    cell.font = Font(bold=True)
        for col_idx, col in enumerate(ws.columns, start=1):
            max_len = max((len(str(cell.value)) for cell in col if cell.value is not None), default=0)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 45)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def write_output_to_bytes(results: dict) -> bytes:
    notes_df = build_export_notes_table()
    partner_df = build_partner_inputs_table(results)
    summary_df = build_summary_table(results)
    unmapped_df = build_unmapped_table(results)
    mapping_df = build_mapping_table(results["config"])
    cleaned_cols = [c for c in results["cleaned_df"].columns if not str(c).startswith("__")]
    norm_cols = [c for c in results["cleaned_df"].columns if str(c).startswith("__")]
    cleaned_export = results["cleaned_df"][cleaned_cols + norm_cols]

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        notes_df.to_excel(writer, sheet_name="Read_me", index=False)
        partner_df.to_excel(writer, sheet_name="Calculation", index=False)
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        unmapped_df.to_excel(writer, sheet_name="Unmapped responses", index=False)
        mapping_df.to_excel(writer, sheet_name="Mapping", index=False)
        cleaned_export.to_excel(writer, sheet_name="Cleaned_Data", index=False)
    buffer.seek(0)
    return format_workbook_bytes(buffer.getvalue())


def dataframes_to_workbook_bytes(sheets: Dict[str, pd.DataFrame]) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for sheet_name, sheet_df in sheets.items():
            sheet_df.to_excel(writer, sheet_name=sheet_name, index=False)
    buffer.seek(0)
    return format_workbook_bytes(buffer.getvalue())


@st.cache_data
def build_template_workbook_bytes() -> bytes:
    template_df = pd.DataFrame(columns=TEMPLATE_COLUMNS)
    overview_df = pd.DataFrame([
        ["Purpose", "Use this workbook structure to prepare an anonymized respondent-level dataset for the calculator."],
        ["One row rule", "Each row should represent one respondent/person."],
        ["Privacy", "Remove names, phone numbers, addresses, ID numbers, case numbers, and free-text identifiers before upload."],
        ["Required tabs", "Fill the Blank dataset sheet, then upload the completed file in the app."],
        ["Optional columns", "disability and donor can be left blank or removed if not used."],
        ["Version", APP_VERSION],
        ["Last updated", APP_LAST_UPDATED],
    ], columns=["Topic", "Guidance"])
    guidance_df = pd.DataFrame([
        ["sex", "Optional", "Use consistent labels such as Female and Male if sex/gender disaggregation is needed."],
        ["age", "Required unless age disaggregation is disabled", "Numeric ages work best. The app can group ages into bands."],
        ["disability", "Optional", "Use consistent yes/no values if disability disaggregation is needed."],
        ["donor", "Optional", "Use if you want to filter results by donor, project, or reporting group."],
        ["SDH. 1", "Required", "Response to KOI question SDH. 1."],
        ["SDH. 2", "Required", "Response to KOI question SDH. 2."],
        ["MEA. 1", "Required", "Response to KOI question MEA. 1."],
        ["MEA. 2", "Required", "Reversed question. In the standard response set, Not really and Not at all are favourable/positive."],
        ["ACC. 1", "Required", "CRM communication means question. Check mapping carefully before calculating."],
        ["ACC. 2", "Required", "Response to KOI question ACC. 2."],
        ["PEM. 1", "Required", "Response to KOI question PEM. 1."],
        ["PEM. 2", "Required", "Response to KOI question PEM. 2."],
    ], columns=["Column", "Required", "Guidance"])
    response_guide_df = pd.DataFrame([
        ["Most questions", "Yes completely | Mostly yes", "Not really | Not at all", "Don't know", "No answer"],
        ["MEA. 2", "Not really | Not at all", "Yes completely | Mostly yes", "Don't know", "No answer"],
        ["ACC. 1", "At least one available CRM communication means known", "CRM communication mean(s) mentioned are not available", "Don't know", "No answer"],
    ], columns=["Question type", "Positive examples", "Negative examples", "Don't know examples", "No answer examples"])
    question_guide_df = pd.DataFrame([
        [QUESTION_LABELS[q], QUESTION_CUES[q], QUESTION_HELP[q]]
        for q in QUESTION_ORDER
    ], columns=["Question", "Short description", "What it means"])
    checklist_df = pd.DataFrame([
        ["Dataset is anonymized", ""],
        ["One row equals one respondent/person", ""],
        ["All 8 KOI question columns are present", ""],
        ["Response labels are consistent across rows", ""],
        ["MEA. 2 mapping has been checked", ""],
        ["ACC. 1 mapping has been checked", ""],
        ["No names, phones, addresses, IDs, case numbers, or identifying free text remain", ""],
    ], columns=["Before upload checklist", "Done"])
    example_df = pd.DataFrame(SAMPLE_DATA[:5], columns=TEMPLATE_COLUMNS)
    return dataframes_to_workbook_bytes({
        "Read_me": overview_df,
        "Blank dataset": template_df,
        "Example anonymized rows": example_df,
        "Column guidance": guidance_df,
        "Question guide": question_guide_df,
        "Response mapping guide": response_guide_df,
        "Before upload checklist": checklist_df,
    })


@st.cache_data
def build_sample_dataset_bytes() -> bytes:
    sample_df = pd.DataFrame(SAMPLE_DATA, columns=TEMPLATE_COLUMNS)
    question_guide_df = pd.DataFrame([
        [QUESTION_LABELS[q], QUESTION_CUES[q], QUESTION_HELP[q]]
        for q in QUESTION_ORDER
    ], columns=["Question", "Short description", "What it means"])
    expected_mapping_df = pd.DataFrame([
        ["SDH. 1", "Yes completely | Mostly yes", "Not really | Not at all", "Don't know", "No answer"],
        ["SDH. 2", "Yes completely | Mostly yes", "Not really | Not at all", "Don't know", "No answer"],
        ["MEA. 1", "Yes completely | Mostly yes", "Not really | Not at all", "Don't know", "No answer"],
        ["MEA. 2", "Not really | Not at all", "Yes completely | Mostly yes", "Don't know", "No answer"],
        ["ACC. 1", "At least one available CRM communication means known", "CRM communication mean(s) mentioned are not available", "Don't know", "No answer"],
        ["ACC. 2", "Yes completely | Mostly yes", "Not really | Not at all", "Don't know", "No answer"],
        ["PEM. 1", "Yes completely | Mostly yes", "Not really | Not at all", "Don't know", "No answer"],
        ["PEM. 2", "Yes completely | Mostly yes", "Not really | Not at all", "Don't know", "No answer"],
    ], columns=["Question", "Positive", "Negative", "Don't know", "No answer"])
    sample_notes_df = pd.DataFrame([
        ["Purpose", "This file is fake anonymized sample data for testing the calculator before uploading real survey data."],
        ["Rows", f"{len(SAMPLE_DATA)} sample respondents across different ages, disability statuses, and programs."],
        ["Expected use", "Upload the Sample dataset sheet and review the suggested response mappings before calculating."],
        ["Privacy", "The sample contains no real people or direct identifiers."],
    ], columns=["Topic", "Note"])
    return dataframes_to_workbook_bytes({
        "Read_me": build_export_notes_table(),
        "Sample notes": sample_notes_df,
        "Sample dataset": sample_df,
        "Question guide": question_guide_df,
        "Expected mapping": expected_mapping_df,
    })


def compact_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", clean_text(value))


def guess_column(columns: List[str], keywords: Iterable[str]) -> Optional[str]:
    scored = []
    for idx, col in enumerate(columns):
        text = clean_text(col)
        compact_text = compact_key(col)
        score = 0
        for keyword in keywords:
            keyword_text = clean_text(keyword)
            keyword_compact = compact_key(keyword)
            if keyword_text and keyword_text in text:
                score += 2
            if keyword_compact and keyword_compact in compact_text:
                score += 3
        if score:
            # preserve left-to-right order as tie breaker
            scored.append((score, -idx, col))
    if not scored:
        return None
    scored.sort(reverse=True)
    return scored[0][2]


def question_column_keywords(question: str) -> List[str]:
    q_code = question.replace("_", "")
    return [
        QUESTION_LABELS[question],
        question.replace("_", " "),
        question.replace("_", "."),
        question.replace("_", ""),
        q_code,
        QUESTION_CUES[question],
        *QUESTION_COLUMN_HINTS[question],
    ]


def make_unique_columns(columns: Iterable[object]) -> List[str]:
    seen: Dict[str, int] = {}
    unique_columns = []
    for column in columns:
        base = str(column).strip() or "Unnamed"
        count = seen.get(base, 0)
        unique_columns.append(base if count == 0 else f"{base} ({count + 1})")
        seen[base] = count + 1
    return unique_columns


def select_column(
    label: str,
    columns: List[str],
    default: Optional[str] = None,
    required: bool = True,
    help_text: Optional[str] = None,
) -> Optional[str]:
    options = columns if required else [""] + columns
    index = options.index(default) if default in options else 0
    value = st.selectbox(label, options=options, index=index, help=help_text)
    return value or None


def render_footer() -> None:
    st.divider()
    st.caption(f"Protection Mainstreaming KOI Calculator version {APP_VERSION} | Last updated {APP_LAST_UPDATED}")


def render_app() -> None:
    st.set_page_config(page_title="Protection Mainstreaming KOI Calculator", layout="wide")
    st.title("Protection Mainstreaming KOI Calculator")
    st.info(
        "Before uploading, anonymize the dataset and remove personally identifiable data "
        "such as names, phone numbers, addresses, ID numbers, or other direct identifiers.",
    )

    with st.expander("Calculation method", expanded=False):
        st.markdown(
            """
- Question result = Positive / (Positive + Negative + Don't know).
- No answer and unmapped or blank responses are excluded from the denominator.
- Final PM KOI is the average of the SDH, MEA, ACC, and PEM element scores.
- MEA. 2 is reversed: for the standard response set, Not really and Not at all are the favourable/positive outcomes.
            """.strip()
        )

    d1, d2 = st.columns(2)
    with d1:
        st.download_button(
            "Download blank template",
            data=build_template_workbook_bytes(),
            file_name="protection_mainstreaming_koi_template.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
    with d2:
        st.download_button(
            "Download sample dataset",
            data=build_sample_dataset_bytes(),
            file_name="protection_mainstreaming_koi_sample.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )

    st.divider()
    uploaded = st.file_uploader("Upload respondent-level Excel dataset", type=["xlsx", "xlsm"])
    if uploaded is None:
        st.info("Upload an Excel dataset to start.")
        render_footer()
        return

    try:
        excel = pd.ExcelFile(uploaded)
    except Exception as exc:
        st.error(f"Could not read the Excel file: {exc}")
        return

    sheet = st.selectbox("Dataset sheet", excel.sheet_names)
    try:
        uploaded.seek(0)
        df = pd.read_excel(uploaded, sheet_name=sheet)
    except Exception as exc:
        st.error(f"Could not read selected sheet: {exc}")
        return

    df.columns = make_unique_columns(df.columns)
    columns = list(df.columns)

    st.divider()
    st.subheader("1. Main column mapping")
    c1, c2, c3, c4 = st.columns(4)
    with c1:
        sex_col = select_column("Optional gender / sex column", columns, guess_column(columns, ["sex", "gender"]), required=False)
    with c2:
        age_disaggregation = st.selectbox("Age disaggregation", ["5-17, 18-49, 50+", "No disaggregation", "Custom bands"], index=0)
    with c3:
        age_col = select_column("Age column", columns, guess_column(columns, ["age", "virst", "varst", "varsta"]), required=(age_disaggregation != "No disaggregation"))
    with c4:
        disability_col = select_column("Optional disability yes/no column", columns, guess_column(columns, ["disability", "disabled", "wgq", "difficulty"]), required=False)

    if age_disaggregation == "Custom bands":
        age_bands = st.text_input("Custom age bands", DEFAULT_AGE_BANDS, help="Format examples: 0-4, 5-17, 18-49, 50+ or 18-59, 60+")
    elif age_disaggregation == "No disaggregation":
        age_bands = ""
    else:
        age_bands = DEFAULT_AGE_BANDS

    donor_col = select_column("Optional donor column. Leave blank to not filter by donor.", columns, guess_column(columns, ["donor", "project", "grant"]), required=False)
    selected_donor_values: List[str] = []
    if donor_col:
        donor_values = unique_display_values(df[donor_col])
        selected_donor_values = st.multiselect(
            "Donor values to include. Leave empty to include all donor values.",
            donor_values,
            default=[],
        )

    st.divider()
    st.subheader("2. Gender and disability value mapping")
    g1, g2, g3 = st.columns(3)
    sex_values = unique_display_values(df[sex_col]) if sex_col else []
    disability_values = unique_display_values(df[disability_col]) if disability_col else []
    with g1:
        female_values = st.multiselect("Values counted as Female", sex_values, default=auto_pick(sex_values, DEFAULT_FEMALE_HINTS))
    with g2:
        male_values = st.multiselect("Values counted as Male", sex_values, default=auto_pick(sex_values, DEFAULT_MALE_HINTS))
    with g3:
        disability_yes_values = st.multiselect("Values counted as disability = Yes", disability_values, default=auto_pick(disability_values, DEFAULT_DISABILITY_YES_HINTS))

    st.divider()
    st.subheader("3. KOI question columns")
    st.caption("Map each KOI question to the matching dataset column.")
    q_cols: Dict[str, str] = {}
    q_layout = st.columns(2)
    for idx, q in enumerate(QUESTION_ORDER):
        with q_layout[idx % 2]:
            q_cols[q] = select_column(
                question_display_label(q),
                columns,
                guess_column(columns, question_column_keywords(q)),
                required=True,
                help_text=QUESTION_HELP[q],
            )

    st.divider()
    st.subheader("4. Response mapping for each KOI question")
    st.caption("Classify each available answer label. Don't know stays in the denominator; No answer, blanks, and unmapped values are excluded.")

    question_mappings: Dict[str, dict] = {}
    for q in QUESTION_ORDER:
        col = q_cols.get(q)
        if not col:
            continue
        values = unique_display_values(df[col])
        with st.expander(f"{question_display_label(q)} response mapping", expanded=True):
            m1, m2, m3, m4 = st.columns(4)
            default_positive = auto_pick(values, DEFAULT_POSITIVE_HINTS)
            default_negative = auto_pick(values, DEFAULT_NEGATIVE_HINTS)
            default_dont_know = auto_pick(values, DEFAULT_DONT_KNOW_HINTS)
            default_no_answer = auto_pick(values, DEFAULT_NO_ANSWER_HINTS)
            if q == "mea_2":
                # MEA.2 is reversed in the standard PM KOI guidance: "Not really"
                # and "Not at all" are counted as the favourable/positive outcome.
                default_positive = auto_pick(values, ["not really", "not at all", "rather no", "mostly no", "no, none", "no none"])
                default_negative = auto_pick(values, ["yes completely", "yes, completely", "completely yes", "mostly yes", "yes mostly", "yes a lot", "yes, a lot", "yes a few", "yes, a few", "a lot", "a few"])
                already_selected = {clean_text(x) for x in default_positive + default_negative + default_dont_know}
                default_no_answer = [v for v in default_no_answer if clean_text(v) not in already_selected]
            if q == "acc_1":
                # Avoid broad hints like "crm", because they can pick both the available
                # and not-available CRM options and create overlapping mappings.
                default_positive = auto_pick(values, ["at least one available crm communication means known", "available crm communication", "available complaint", "yes completely", "yes, completely", "mostly yes"])
                default_negative = auto_pick(values, ["crm communication mean(s) mentioned are not available", "mentioned are not available", "not really", "not at all"])
                already_selected = {clean_text(x) for x in default_positive + default_negative + default_dont_know}
                default_no_answer = [v for v in default_no_answer if clean_text(v) not in already_selected]
            with m1:
                positive = st.multiselect("Positive", values, default=default_positive, key=f"{q}_positive")
            with m2:
                negative = st.multiselect("Negative", values, default=default_negative, key=f"{q}_negative")
            with m3:
                dont_know = st.multiselect("Don't know", values, default=default_dont_know, key=f"{q}_dont_know")
            with m4:
                no_answer = st.multiselect("No answer", values, default=default_no_answer, key=f"{q}_no_answer")
            question_mappings[q] = {"positive": positive, "negative": negative, "dont_know": dont_know, "no_answer": no_answer}

    config = {
        "columns": {"sex": sex_col, "age": age_col, "disability": disability_col, "donor": donor_col, **q_cols},
        "age_disaggregation": age_disaggregation,
        "age_bands": age_bands,
        "donor_values": selected_donor_values,
        "sex_values": {"female": female_values, "male": male_values},
        "disability_yes_values": disability_yes_values,
        "question_mappings": question_mappings,
    }

    st.subheader("5. Review mapping before calculation")
    review_df = apply_donor_filter(df, donor_col, selected_donor_values) if donor_col else df
    st.dataframe(build_mapping_review_table(review_df, config), use_container_width=True)
    errors, warnings = build_preflight_checks(review_df, config)
    if errors:
        st.error("Fix these before calculating:\n\n" + "\n".join(f"- {error}" for error in errors))
    if warnings:
        st.warning("Review these before final reporting:\n\n" + "\n".join(f"- {warning}" for warning in warnings))

    with st.expander("Preview uploaded data", expanded=False):
        st.dataframe(df.head(50), use_container_width=True)

    calculate_clicked = st.button("Calculate KOI", type="primary", disabled=bool(errors))
    if errors or not calculate_clicked:
        render_footer()
        return

    try:
        results = calculate_koi(df, config)
    except Exception as exc:
        st.error(str(exc))
        return

    final = results["final_scores"].get("total")
    r1, r2, r3, r4 = st.columns(4)
    r1.metric("Original rows", results["original_rows"])
    r2.metric("Rows after donor filter", results["filtered_rows"])
    r3.metric("Final PM KOI", "N/A" if final is None else f"{final:.2%}")
    negative_unique = results.get("unique_people_with_negative")
    r4.metric(NEGATIVE_PERSON_METRIC_LABEL, negative_unique)
    st.caption(
        "This is a diagnostic count, not the inverse of the final PM KOI score. The final score averages "
        "question percentages across SDH, MEA, ACC, and PEM; this count flags each row/person once if they "
        "have any negative KOI answer. MEA. 2 uses reversed scoring."
    )

    summary_df = build_summary_table(results)
    partner_df = build_partner_inputs_table(results)
    unmapped_df = build_unmapped_table(results)

    st.subheader("Summary")
    st.caption("Question results use Positive / (Positive + Negative + Don't know). No answer and unmapped/blank values are excluded.")
    st.dataframe(summary_df, use_container_width=True)

    st.subheader("Calculation")
    st.dataframe(partner_df, use_container_width=True)

    if not unmapped_df.empty and unmapped_df.iloc[0, 0] != "No unselected responses found":
        st.warning("Some values were not classified as Positive, Negative, Don't know, or No answer. They are now treated as No answer and excluded from the denominator. Review them before final reporting.")
    st.subheader("Unmapped responses excluded from denominator")
    st.dataframe(unmapped_df, use_container_width=True)

    output = write_output_to_bytes(results)
    st.download_button(
        "Download Excel results",
        data=output,
        file_name="protection_mainstreaming_koi_results.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    render_footer()


if __name__ == "__main__":
    render_app()
